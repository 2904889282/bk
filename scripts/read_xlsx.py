import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook

path = os.path.join(os.path.expanduser('~/Desktop'), '新建 XLSX 工作表.xlsx')
wb = load_workbook(path)
print('SHEETS:', wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f'\n{"="*80}')
    print(f'Sheet: {sname}  ({ws.max_row} rows x {ws.max_column} cols)')
    print('='*80)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        vals = [str(c.value).replace('\n','↙') if c.value is not None else '∅' for c in row]
        print(' | '.join(vals))
